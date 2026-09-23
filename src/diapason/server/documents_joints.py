"""Les documents joints à un message du chat — lus, bornés, et annoncés.

22 septembre 2026. Un extracteur existait dans ``server/upload_router.py``
depuis longtemps, mais il verse dans le CORPUS de la recherche approfondie :
une bibliothèque qu'on interroge, pas une pièce qu'on met sous les yeux.
« Regarde ce PDF » et « ajoute ce PDF à ta bibliothèque » sont deux gestes
différents, et seul le second existait. Ses dépendances n'étaient d'ailleurs
pas installées — ``make setup`` ne pose pas ``memory-pdf`` — donc un PDF
joint rendait une erreur 500.

Ce module lit le document et le borne. Le bornage n'est pas un détail : un
PDF de cent pages fait cinquante mille mots, la fenêtre du modèle du
quotidien en tient trente-deux mille jetons, et un texte tronqué en silence
fait répondre avec assurance sur une moitié de document. **Ce qui est coupé
est dit** — au modèle dans le texte, et à l'usager dans le compte rendu
(§5 : ne jamais faire semblant).
"""

from __future__ import annotations

import base64
import binascii
import io
import logging
import re
from dataclasses import dataclass

logger = logging.getLogger(__name__)

# Le fichier lui-même. Un mémoire de maîtrise en PDF dépasse rarement ça, et
# au-delà on transporte des octets qu'on va de toute façon couper.
TAILLE_MAX = 10 * 1024 * 1024
# Le texte retenu. Un caractère vaut grossièrement un quart de jeton en
# français : 48 000 caractères pèsent donc ~12 000 jetons, soit un bon tiers
# de la fenêtre de 32 Ko du 9b, ce qui laisse la place à la conversation et
# à la réponse.
CARACTERES_MAX = 48_000
# Deux documents suffisent pour comparer ; trois remplissent la fenêtre.
NOMBRE_MAX = 2

_EXTENSIONS = {".txt", ".md", ".csv", ".pdf", ".docx"}
_SIGNATURES: tuple[tuple[bytes, str], ...] = (
    (b"%PDF-", ".pdf"),
    # .docx est un ZIP ; la distinction d'avec un .zip quelconque se fait à
    # l'ouverture, pas sur la signature.
    (b"PK\x03\x04", ".docx"),
)
# Les blancs qu'un extracteur de PDF sème entre les colonnes et les en-têtes.
_BLANCS = re.compile(r"[ \t]{2,}")
_LIGNES_VIDES = re.compile(r"\n{3,}")


class DocumentRefuse(ValueError):
    """Ce qui est arrivé ne peut pas être lu, ou ne devrait pas l'être."""


@dataclass(frozen=True)
class DocumentLu:
    """Un document prêt à être montré au modèle, et ce qu'on en dit."""

    nom: str
    texte: str
    caracteres: int
    """La longueur du texte RETENU, pas celle du document."""
    caracteres_source: int
    """Ce que le document contenait avant le bornage."""
    pages: int | None
    """Le nombre de pages, pour un PDF ; None pour les autres formats."""

    @property
    def tronque(self) -> bool:
        return self.caracteres < self.caracteres_source


def _extension(nom: str, octets: bytes) -> str:
    """L'extension retenue : celle du nom, confirmée par les octets quand
    ils parlent.

    Un nom de fichier est une déclaration du client. Pour un PDF les octets
    tranchent ; pour du texte brut il n'y a rien à lire, et on fait avec.
    """
    depuis_le_nom = ("." + nom.rsplit(".", 1)[-1].lower()) if "." in nom else ""
    for signature, ext in _SIGNATURES:
        if octets.startswith(signature):
            # Un .docx annoncé .pdf, ou l'inverse : les octets gagnent.
            return ext
    if depuis_le_nom in _EXTENSIONS:
        if octets[:5] == b"%PDF-" or octets[:4] == b"PK\x03\x04":
            return depuis_le_nom  # déjà traité plus haut, par sûreté
        return depuis_le_nom
    raise DocumentRefuse(
        f"{nom} : format non pris en charge. "
        "Les documents acceptés sont .txt, .md, .csv, .pdf et .docx."
    )


def _texte_du_pdf(octets: bytes) -> tuple[str, int]:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover - dépendance déclarée
        raise DocumentRefuse(
            "La lecture des PDF n'est pas installée sur ce serveur "
            "(extra « documents »)."
        ) from exc
    try:
        with pdfplumber.open(io.BytesIO(octets)) as pdf:
            pages = [p.extract_text() or "" for p in pdf.pages]
    except Exception as exc:  # noqa: BLE001 - un PDF cassé n'est pas une panne
        raise DocumentRefuse(f"PDF illisible : {exc}") from exc
    return "\n\n".join(pages), len(pages)


def _texte_du_docx(octets: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - dépendance déclarée
        raise DocumentRefuse(
            "La lecture des .docx n'est pas installée sur ce serveur "
            "(extra « documents »)."
        ) from exc
    try:
        doc = Document(io.BytesIO(octets))
    except Exception as exc:  # noqa: BLE001
        raise DocumentRefuse(f"Document Word illisible : {exc}") from exc
    morceaux = [p.text for p in doc.paragraphs if p.text.strip()]
    # Les tableaux portent souvent l'essentiel d'un document Word, et les
    # ignorer donnait un texte qui semblait complet tout en ayant perdu les
    # chiffres.
    for table in doc.tables:
        for ligne in table.rows:
            cellules = [c.text.strip() for c in ligne.cells if c.text.strip()]
            if cellules:
                morceaux.append(" | ".join(cellules))
    return "\n\n".join(morceaux)


def _nettoyer(texte: str) -> str:
    """Les blancs d'un extracteur ne portent aucun sens et coûtent des jetons."""
    texte = _BLANCS.sub(" ", texte)
    texte = _LIGNES_VIDES.sub("\n\n", texte)
    return texte.strip()


def lire(nom: str, contenu_base64: str) -> DocumentLu:
    """Lit un document joint, ou dit pourquoi il ne peut pas l'être."""
    brut = (contenu_base64 or "").strip()
    if brut.startswith("data:"):
        brut = brut.partition(",")[2]
    if not brut:
        raise DocumentRefuse(f"{nom} : document vide.")
    try:
        octets = base64.b64decode(brut, validate=True)
    except (binascii.Error, ValueError) as exc:
        raise DocumentRefuse(f"{nom} : contenu illisible.") from exc
    if len(octets) > TAILLE_MAX:
        raise DocumentRefuse(
            f"{nom} : {len(octets) // 1024 // 1024} Mo, "
            f"maximum {TAILLE_MAX // 1024 // 1024} Mo."
        )

    ext = _extension(nom, octets)
    pages: int | None = None
    if ext == ".pdf":
        texte, pages = _texte_du_pdf(octets)
    elif ext == ".docx":
        texte = _texte_du_docx(octets)
    else:
        try:
            texte = octets.decode("utf-8")
        except UnicodeDecodeError:
            texte = octets.decode("latin-1", errors="replace")

    texte = _nettoyer(texte)
    if not texte:
        raise DocumentRefuse(
            f"{nom} : aucun texte lisible. Un PDF de pages scannées est une "
            "image ; il faudrait le passer par la reconnaissance de caractères."
        )
    source = len(texte)
    if source > CARACTERES_MAX:
        # Coupé sur une frontière de mot : une phrase tronquée au milieu d'un
        # mot se lit comme une donnée corrompue.
        texte = texte[:CARACTERES_MAX].rsplit(" ", 1)[0]
    return DocumentLu(
        nom=nom,
        texte=texte,
        caracteres=len(texte),
        caracteres_source=source,
        pages=pages,
    )


def en_texte(document: DocumentLu) -> str:
    """Le bloc montré au modèle : le document, nommé, et sa coupure dite.

    Le modèle doit savoir qu'il ne voit qu'un début, sinon il répond « le
    document ne mentionne pas X » avec l'assurance de qui a tout lu.
    """
    entete = f"--- Document joint : {document.nom}"
    if document.pages:
        entete += f" ({document.pages} page{'s' if document.pages > 1 else ''})"
    entete += " ---"
    pied = ""
    if document.tronque:
        lus = round(100 * document.caracteres / document.caracteres_source)
        pied = (
            f"\n\n--- Fin de l'extrait : seuls les {lus} % du début de "
            f"{document.nom} sont montrés ici. Ne conclus pas sur ce qui "
            "manque ; dis que le reste n'a pas été lu. ---"
        )
    return f"{entete}\n{document.texte}{pied}"


def composer(contenu: str, documents: list[dict] | None) -> str:
    """Le message tel que le modèle le lit : les documents, puis la question.

    Les documents sont extraits UNE fois, quand l'usager les joint (route
    ``/v1/chat/documents``), et le message ne porte ensuite que leur texte —
    relire un PDF à chaque tour de la conversation coûterait une seconde par
    tour pour un résultat identique.

    Les blocs passent AVANT la question : « voici la pièce, voici ce que j'en
    demande » est l'ordre dans lequel on lit, et celui dans lequel un modèle
    répond le mieux.
    """
    if not documents:
        return contenu
    blocs = []
    for d in documents:
        nom = str(d.get("nom") or "document")
        texte = str(d.get("texte") or "").strip()
        if not texte:
            continue
        pages = d.get("pages")
        entete = f"--- Document joint : {nom}"
        if isinstance(pages, int) and pages > 0:
            entete += f" ({pages} page{'s' if pages > 1 else ''})"
        entete += " ---"
        bloc = f"{entete}\n{texte}"
        if d.get("tronque"):
            bloc += (
                f"\n\n--- Fin de l'extrait : seul le début de {nom} est montré "
                "ici. Ne conclus pas sur ce qui manque ; dis que le reste n'a "
                "pas été lu. ---"
            )
        blocs.append(bloc)
    if not blocs:
        return contenu
    return "\n\n".join(blocs) + ("\n\n" + contenu if contenu.strip() else "")


def lire_tous(joints: list[dict] | None) -> list[DocumentLu]:
    """Les documents d'un message. Un seul refusé refuse le message entier —
    livrer les autres en silence ferait croire que celui-là a été lu (§5)."""
    if not joints:
        return []
    if len(joints) > NOMBRE_MAX:
        raise DocumentRefuse(
            f"{len(joints)} documents joints, maximum {NOMBRE_MAX} par message."
        )
    return [
        lire(str(d.get("nom") or "document"), str(d.get("contenu") or ""))
        for d in joints
    ]


__all__ = [
    "CARACTERES_MAX",
    "composer",
    "NOMBRE_MAX",
    "TAILLE_MAX",
    "DocumentLu",
    "DocumentRefuse",
    "en_texte",
    "lire",
    "lire_tous",
]
