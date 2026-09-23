"""Les documents joints à un message du chat.

22/09/2026 : un extracteur existait, mais il versait dans le CORPUS de la
recherche approfondie — une bibliothèque qu'on interroge, pas une pièce
qu'on met sous les yeux. Et aucune de ses dépendances n'était installée :
`make setup` ne pose pas `memory-pdf`, donc un PDF joint rendait un 500.
"""

import base64
import io

import pytest

from diapason.server.documents_joints import (
    CARACTERES_MAX,
    NOMBRE_MAX,
    DocumentRefuse,
    composer,
    lire,
    lire_tous,
)

pdfplumber = pytest.importorskip("pdfplumber")
docx = pytest.importorskip("docx")


def en_base64(octets: bytes) -> str:
    return base64.b64encode(octets).decode()


def un_docx(paragraphes, tableau=None) -> bytes:
    from docx import Document

    d = Document()
    for p in paragraphes:
        d.add_paragraph(p)
    if tableau:
        t = d.add_table(rows=len(tableau), cols=len(tableau[0]))
        for i, ligne in enumerate(tableau):
            for j, cellule in enumerate(ligne):
                t.rows[i].cells[j].text = cellule
    tampon = io.BytesIO()
    d.save(tampon)
    return tampon.getvalue()


def un_pdf(phrase: str) -> bytes:
    """Un PDF minimal écrit à la main : pas de dépendance d'écriture."""
    flux = f"BT /F1 14 Tf 72 720 Td ({phrase}) Tj ET".encode()
    objets = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length "
        + str(len(flux)).encode()
        + b" >>\nstream\n"
        + flux
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    sortie, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, o in enumerate(objets, 1):
        offsets.append(len(sortie))
        sortie += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
        debut = len(sortie)
    sortie += f"xref\n0 {len(objets) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        sortie += f"{off:010d} 00000 n \n".encode()
    sortie += (
        f"trailer\n<< /Size {len(objets) + 1} /Root 1 0 R >>\n"
        f"startxref\n{debut}\n%%EOF\n".encode()
    )
    return bytes(sortie)


class TestCeQuiSeLit:
    def test_un_pdf_rend_son_texte_et_son_nombre_de_pages(self):
        d = lire(
            "bail.pdf", en_base64(un_pdf("Le loyer mensuel est de 1 450 dollars."))
        )
        assert "1 450 dollars" in d.texte
        assert d.pages == 1
        assert not d.tronque

    def test_un_docx_rend_ses_paragraphes(self):
        d = lire("note.docx", en_base64(un_docx(["Première ligne.", "Seconde ligne."])))
        assert "Première ligne." in d.texte and "Seconde ligne." in d.texte
        assert d.pages is None, "un .docx n'a pas de pages avant d'être mis en page"

    def test_les_tableaux_d_un_docx_ne_sont_pas_perdus(self):
        """Ils portent souvent l'essentiel — les ignorer donnait un texte qui
        semblait complet tout en ayant perdu les chiffres."""
        octets = un_docx(
            ["Rapport."], tableau=[["Trimestre", "Montant"], ["T3", "148 200 $"]]
        )
        d = lire("rapport.docx", en_base64(octets))
        assert "148 200 $" in d.texte
        assert "Trimestre | Montant" in d.texte

    def test_le_texte_brut_passe_tel_quel(self):
        d = lire("notes.md", en_base64("# Titre\n\nUn paragraphe.".encode()))
        assert d.texte == "# Titre\n\nUn paragraphe."

    def test_l_entete_data_est_acceptee(self):
        nu = en_base64(b"bonjour")
        assert lire("a.txt", f"data:text/plain;base64,{nu}").texte == "bonjour"

    def test_les_blancs_d_extraction_sont_resserres(self):
        """Ils ne portent aucun sens et coûtent des jetons."""
        d = lire(
            "a.txt", en_base64("colonne     suivante\n\n\n\n\nparagraphe".encode())
        )
        assert d.texte == "colonne suivante\n\nparagraphe"


class TestCeQuiEstRefuse:
    def test_un_format_inconnu(self):
        with pytest.raises(DocumentRefuse, match="non pris en charge"):
            lire("archive.zip", en_base64(b"\x00\x01\x02\x03 pas un document"))

    def test_un_document_vide(self):
        with pytest.raises(DocumentRefuse, match="vide"):
            lire("a.pdf", "")

    def test_un_pdf_cassé_n_est_pas_une_panne(self):
        with pytest.raises(DocumentRefuse, match="illisible"):
            lire("casse.pdf", en_base64("%PDF-1.4 mais rien derrière".encode()))

    def test_un_pdf_sans_texte_le_dit(self):
        """Un PDF de pages scannées est une image : le silence ferait croire
        que le document ne contenait rien."""
        vide = un_pdf("")
        with pytest.raises(DocumentRefuse, match="reconnaissance de caractères"):
            lire("scan.pdf", en_base64(vide))

    def test_trop_de_documents(self):
        joints = [
            {"nom": f"d{i}.txt", "contenu": en_base64(b"x")}
            for i in range(NOMBRE_MAX + 1)
        ]
        with pytest.raises(DocumentRefuse, match=f"maximum {NOMBRE_MAX}"):
            lire_tous(joints)

    def test_un_seul_refuse_refuse_le_message(self):
        with pytest.raises(DocumentRefuse):
            lire_tous(
                [
                    {"nom": "bon.txt", "contenu": en_base64(b"texte")},
                    {"nom": "mauvais.zip", "contenu": en_base64(b"\x00\x01\x02\x03")},
                ]
            )


class TestLaCoupureSeDit:
    """Un texte tronqué en silence fait répondre avec assurance sur une
    moitié de document."""

    def test_au_dela_de_la_borne_le_texte_est_coupe(self):
        long = ("mot " * (CARACTERES_MAX // 2)).encode()
        d = lire("gros.txt", en_base64(long))
        assert d.tronque
        assert d.caracteres <= CARACTERES_MAX
        assert d.caracteres_source > CARACTERES_MAX
        assert not d.texte.endswith("mo"), "la coupe tombe sur une frontière de mot"

    def test_en_deca_rien_n_est_coupe(self):
        d = lire("petit.txt", en_base64(b"court"))
        assert not d.tronque and d.caracteres == d.caracteres_source


class TestCeQueLeModeleLit:
    def test_le_document_precede_la_question(self):
        sortie = composer(
            "Quel est le loyer ?",
            [{"nom": "bail.pdf", "texte": "Le loyer est de 1 450 $.", "pages": 3}],
        )
        assert sortie.index("bail.pdf") < sortie.index("Quel est le loyer ?")
        assert "(3 pages)" in sortie

    def test_la_coupure_est_dite_au_modele(self):
        sortie = composer(
            "Résume.",
            [{"nom": "these.pdf", "texte": "début…", "pages": 200, "tronque": True}],
        )
        assert "seul le début de these.pdf est montré" in sortie
        assert "Ne conclus pas sur ce qui manque" in sortie

    def test_sans_document_le_message_ne_change_pas(self):
        assert composer("bonjour", None) == "bonjour"
        assert composer("bonjour", []) == "bonjour"

    def test_un_document_sans_texte_est_ignore(self):
        assert composer("bonjour", [{"nom": "vide.txt", "texte": "  "}]) == "bonjour"

    def test_un_document_sans_question_reste_seul(self):
        sortie = composer("", [{"nom": "a.txt", "texte": "contenu"}])
        assert sortie.endswith("contenu")
        assert not sortie.endswith("\n\n")

    def test_deux_documents_gardent_leur_ordre(self):
        sortie = composer(
            "Compare-les.",
            [{"nom": "a.txt", "texte": "premier"}, {"nom": "b.txt", "texte": "second"}],
        )
        assert (
            sortie.index("a.txt") < sortie.index("b.txt") < sortie.index("Compare-les.")
        )


class TestLeCheminComplet:
    def test_la_conversion_compose_les_documents(self):
        from diapason.server.models import ChatMessage
        from diapason.server.routes import _to_messages

        (message,) = _to_messages(
            [
                ChatMessage(
                    role="user",
                    content="Combien ?",
                    documents=[{"nom": "f.txt", "texte": "Le montant est 42 $."}],
                )
            ]
        )
        assert "Le montant est 42 $." in message.content
        assert message.content.index("f.txt") < message.content.index("Combien ?")

    def test_un_message_sans_document_traverse_intact(self):
        from diapason.server.models import ChatMessage
        from diapason.server.routes import _to_messages

        (message,) = _to_messages([ChatMessage(role="user", content="bonjour")])
        assert message.content == "bonjour"
